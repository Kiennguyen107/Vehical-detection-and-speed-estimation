import sys
import cv2
import numpy as np
import pandas as pd
from pathlib import Path
from datetime import datetime
from queue import Queue, Empty

from PyQt6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
    QPushButton, QLabel, QFileDialog, QTextEdit, QTabWidget,
    QGroupBox, QTableWidget, QTableWidgetItem, QFrame,
    QSplitter, QHeaderView, QSizePolicy
)
from PyQt6.QtCore import Qt, QThread, pyqtSignal
from PyQt6.QtGui import QPixmap, QImage, QFont

from matplotlib.backends.backend_qt5agg import FigureCanvasQTAgg
from matplotlib.figure import Figure

from ultralytics import YOLO
from paddleocr import PaddleOCR

from database import (
    init_db, insert_speed_violation, insert_distance_violation,
    update_plate, update_color, update_dist_violation,
    update_dist_plate
)

from lane_detection import (
    detect_lanes,
    draw_lanes_on_frame,
    get_lane_of_point,
    filter_coeffs_in_roi,
)
from speed_estimation import (
    SpeedEstimator,
    DEFAULT_ROI,
)
from car_color_detection import detect_car_color

# ===== PLATE NORMALIZATION =====
DIGIT_TO_CHAR = {
    '0': 'O', '1': 'I', '8': 'B', '6': 'G', '5': 'S',
}
CHAR_TO_DIGIT = {
    'O': '0', 'I': '1', 'B': '8', 'G': '6',
    'S': '5', 'J': '3', 'Z': '2', 'T': '7',
}

def normalize_plate(text: str) -> str:
    text = text.replace(" ", "").upper()
    text = text.replace("-", "").replace(".", "")
    if len(text) < 3:
        return text
    result = list(text)
    for i in range(0, 2):
        if result[i].isalpha():
            result[i] = CHAR_TO_DIGIT.get(result[i], result[i])
    if result[2].isdigit():
        result[2] = DIGIT_TO_CHAR.get(result[2], result[2])
    for i in range(3, len(result)):
        if result[i].isalpha():
            result[i] = CHAR_TO_DIGIT.get(result[i], result[i])
    return "".join(result)

# ===== CẤU HÌNH =====
VIDEO_PATH        = "D:/DATN/video/IMG_0556.MOV"
LANE_MODEL_PATH   = "D:/DATN/model/marrking/last_ver9.pt"
CAR_MODEL_PATH    = "C:/Users/kienn/Downloads/best_detect_car_v3.pt"
PLATE_MODEL_PATH  = "D:/DATN/model/best_result_plate_ver1.pt"
CROPS_OUTPUT_DIR  = "roi_crops"
EXCEL_OUTPUT_DIR  = "violations_excel"

LANE_SPEED_LIMITS = {
    1: {"min": 0, "max": 61},
    2: {"min": 0, "max": 71},
    3: {"min": 0, "max": 81},
    4: {"min": 0, "max": 81},
}

LANE_SPEED_LIMITS_UI = {
    1: {"min": 0, "max": 60},
    2: {"min": 0, "max": 70},
    3: {"min": 0, "max": 80},
    4: {"min": 0, "max": 80},
}

COLOR_ROI_CONFIG = {
    'y_top':   0.60,
    'y_bot':   0.88,
    'x_left':  0.20,
    'x_right': 0.80,
}
COLOR_N_CLUSTERS = 16


# ===== PLATE OCR THREAD =====
class PlateOCRThread(QThread):
    plate_detected = pyqtSignal(int, str)
    
    def __init__(self, plate_model_path: str, queue: Queue):
        super().__init__()
        self.queue          = queue
        self.is_running     = False
        self.plate_crop_dir = "roi_crops_plates"
        Path(self.plate_crop_dir).mkdir(parents=True, exist_ok=True)
        self.plate_model = YOLO(plate_model_path)
        self.ocr         = PaddleOCR(
            lang='en',
            use_doc_orientation_classify=False,
            use_doc_unwarping=False,
            use_textline_orientation=True,
        )

    def run(self):
        self.is_running = True
        while self.is_running:
            try:
                item = self.queue.get(timeout=1)
            except Empty:
                continue
            tid        = item['tid']
            crop_path  = item['path']
            plate_text = self._process(crop_path)
            self.plate_detected.emit(tid, plate_text)

    def _process(self, crop_path: str) -> str:
        try:
            results = self.plate_model(crop_path, verbose=False)
            boxes   = results[0].boxes
            if boxes is None or len(boxes) == 0:
                return "N/A"
            img         = results[0].orig_img
            plate_crops = []
            for i, box in enumerate(boxes):
                x1, y1, x2, y2 = map(int, box.xyxy[0])
                cropped = img[y1:y2, x1:x2]
                if cropped.size == 0:
                    continue
                fname    = Path(crop_path).stem
                tmp_path = f"{self.plate_crop_dir}/{fname}_plate_{i}.jpg"
                cv2.imwrite(tmp_path, cropped)
                plate_crops.append(tmp_path)
            if not plate_crops:
                return "N/A"
            texts = []
            for plate_path in plate_crops:
                ocr_result = self.ocr.predict(plate_path)
                for res in ocr_result:
                    rec_texts = res['rec_texts'] if isinstance(res, dict) else res.get('rec_texts', [])
                    texts.extend(rec_texts)
            plate_text = " ".join(texts).strip()
            if plate_text:
                plate_text = normalize_plate(plate_text)
            return plate_text if plate_text else "N/A"
        except Exception as e:
            print(f"  PlateOCR error: {e}")
            return "N/A"

    def stop(self):
        self.is_running = False


# ===== COLOR DETECT THREAD =====
class ColorDetectThread(QThread):
    color_detected = pyqtSignal(int, str)

    def __init__(self, queue: Queue):
        super().__init__()
        self.queue      = queue
        self.is_running = False

    def run(self):
        self.is_running = True
        while self.is_running:
            try:
                item = self.queue.get(timeout=1)
            except Empty:
                continue
            tid       = item['tid']
            crop_path = item['path']
            color     = self._process(crop_path)
            self.color_detected.emit(tid, color)

    def _process(self, crop_path: str) -> str:
        try:
            img = cv2.imread(crop_path)
            if img is None:
                return "N/A"
            color_name, _, _ = detect_car_color(
                img,
                n_clusters=COLOR_N_CLUSTERS,
                roi_config=COLOR_ROI_CONFIG,
                debug=False,
            )
            return color_name if color_name else "unknown"
        except Exception as e:
            print(f"  ColorDetect error: {e}")
            return "N/A"

    def stop(self):
        self.is_running = False


# ===== VIDEO PROCESSING THREAD =====
class VideoThread(QThread):
    frame_ready          = pyqtSignal(np.ndarray)
    stats_ready          = pyqtSignal(dict)
    violation_added      = pyqtSignal(dict)
    dist_violation_added = pyqtSignal(dict)
    finished             = pyqtSignal()

    def __init__(self):
        super().__init__()
        self.video_path    = None
        self.car_model     = None
        self.estimator     = None
        self.solid_coeffs  = None
        self.dotted_coeffs = None
        self.all_lines     = None
        self.vp            = None
        self.roi_pts       = None
        self.is_running    = False

    def setup(self, video_path, car_model, estimator,
              solid_coeffs, dotted_coeffs, all_lines, vp, roi_pts):
        self.video_path    = video_path
        self.car_model     = car_model
        self.estimator     = estimator
        self.solid_coeffs  = solid_coeffs
        self.dotted_coeffs = dotted_coeffs
        self.all_lines     = all_lines
        self.vp            = vp
        self.roi_pts       = roi_pts

    def run(self):
        cap = cv2.VideoCapture(self.video_path)
        if not cap.isOpened():
            return

        self.is_running      = True
        cap.set(cv2.CAP_PROP_POS_FRAMES, 0)
        prev_violations      = 0
        prev_dist_violations = 0

        import time
        t0 = time.time()
        frame_count = 0

        # Cache lane overlay 1 lần duy nhất
        ret, first = cap.read()
        lane_overlay = draw_lanes_on_frame(
            first, self.solid_coeffs, self.dotted_coeffs,
            self.all_lines, self.vp, roi_pts=self.roi_pts
        )
        # Tính phần diff giữa lane_overlay và frame gốc — chỉ lưu phần lane
        lane_diff = cv2.subtract(lane_overlay, first)
        cap.set(cv2.CAP_PROP_POS_FRAMES, 0)  # reset về đầu

        while self.is_running and cap.isOpened():
            ret, frame = cap.read()
            if not ret:
                break

            frame_count += 1
            if frame_count % 30 == 0:
                elapsed = time.time() - t0
                # print(f"FPS thực tế: {frame_count / elapsed:.1f}")

            # Blend lane diff lên frame mới thay vì gọi draw_lanes_on_frame
            frame_with_lanes = cv2.add(frame, lane_diff)

            annotated, stats = self.estimator.process_frame(
                frame_with_lanes, frame,
                self.car_model, self.all_lines, get_lane_of_point
            )

            self.frame_ready.emit(annotated)
            self.stats_ready.emit(stats)

            if stats['violations_count'] > prev_violations:
                for v in self.estimator.violations[prev_violations:]:
                    self.violation_added.emit(v)
                prev_violations = stats['violations_count']

            n_dist = len(self.estimator.distance_violations)
            if n_dist > prev_dist_violations:
                for dv in self.estimator.distance_violations[prev_dist_violations:]:
                    self.dist_violation_added.emit(dv)
                prev_dist_violations = n_dist

        cap.release()
        self.finished.emit()

    def stop(self):
        self.is_running = False


# ===== MATPLOTLIB CHART =====
class MplCanvas(FigureCanvasQTAgg):
    def __init__(self, width=5, height=4):
        fig       = Figure(figsize=(width, height), dpi=100)
        self.axes = fig.add_subplot(111)
        super().__init__(fig)
        fig.subplots_adjust(top=0.88, bottom=0.08, left=0.08, right=0.95)


# ===== CLICKABLE LABEL =====
class ClickableLabel(QLabel):
    clicked = pyqtSignal(int, int)

    def mousePressEvent(self, event):
        if event.button() == Qt.MouseButton.LeftButton:
            self.clicked.emit(event.pos().x(), event.pos().y())


# ===== MAIN WINDOW =====
class TrafficApp(QMainWindow):

    # Column indices cho vtable (Speed Violations) — không có cột ID
    COL_TIME  = 0
    COL_DATE  = 1
    COL_SPEED = 2
    COL_OVER  = 3
    COL_PLATE = 4
    COL_COLOR = 5
    COL_DIST  = 6

    def __init__(self):
        super().__init__()
        self.setWindowTitle(" Traffic Speed Violation Detection System")
        self.setGeometry(100, 50, 1600, 950)
        
        # Database
        init_db()
        
        # State
        self.video_path = VIDEO_PATH
        self.first_frame = None
        self.fps = 30
        self.lane_model = None
        self.car_model = None
        self.solid_coeffs = None
        self.dotted_coeffs = None
        self.all_lines  = None
        self.vp = None
        self.roi_pts = []
        self.estimator = None
        self.tid_to_row = {}   # tid → row index trong vtable
        self.tid_to_frame_path = {}   # tid → đường dẫn full frame có bbox
        self.tid_to_dist_row = {}
        self.tid_to_dist_frame_path= {}

        # Queues & threads
        self.plate_queue = Queue()
        self.color_queue = Queue()
        self.plate_thread = None
        self.color_thread = None

        self.thread = VideoThread()
        self.thread.frame_ready.connect(self._on_frame)
        self.thread.stats_ready.connect(self._on_stats)
        self.thread.violation_added.connect(self._on_violation)
        self.thread.dist_violation_added.connect(self._on_dist_violation)
        self.thread.finished.connect(self._on_finished)

        self._build_ui()

    # ===== UI BUILD =====
    def _build_ui(self):
        central = QWidget()
        self.setCentralWidget(central)
        root = QVBoxLayout(central)
        root.addWidget(self._make_header())
        root.addWidget(self._make_metrics())
        tabs = QTabWidget()
        tabs.addTab(self._make_setup_tab(),      " 1. Setup")
        tabs.addTab(self._make_detection_tab(),  " 2. Detection")
        tabs.addTab(self._make_violations_tab(), " 3. Violations")
        root.addWidget(tabs)

    def _make_header(self):
        box = QGroupBox()
        box.setStyleSheet(
            "QGroupBox { background: qlineargradient(x1:0,y1:0,x2:1,y2:0,"
            "stop:0 #2c3e50, stop:1 #3498db); border-radius:5px; padding:20px; }"
        )
        lay = QVBoxLayout()
        t   = QLabel(" Traffic Speed Violation Detection System")
        t.setFont(QFont("Arial", 22, QFont.Weight.Bold))
        t.setStyleSheet("color:white;")
        t.setAlignment(Qt.AlignmentFlag.AlignCenter)
        s   = QLabel("Desktop Application — Real-time Performance")
        s.setStyleSheet("color:#ecf0f1;")
        s.setAlignment(Qt.AlignmentFlag.AlignCenter)
        lay.addWidget(t); lay.addWidget(s)
        box.setLayout(lay)
        return box

    def _make_metrics(self):
        panel = QFrame()
        lay   = QHBoxLayout()
        self.m_vehicles   = self._metric_box("Total Vehicles",   "0",    "#2ecc71")
        self.m_violations = self._metric_box("Speed Violations",  "0",    "#e74c3c")
        self.m_rate       = self._metric_box("Violation Rate",   "0.0%", "#f39c12")
        self.m_tracking   = self._metric_box("Tracking",          "0",    "#3498db")
        for m in [self.m_vehicles, self.m_violations, self.m_rate, self.m_tracking]:
            lay.addWidget(m)
        panel.setLayout(lay)
        return panel

    def _metric_box(self, title, value, color):
        box = QGroupBox()
        box.setStyleSheet(
            f"QGroupBox {{ background-color:{color}; border-radius:8px; padding:15px; }}"
        )
        lay   = QVBoxLayout()
        lbl_t = QLabel(title)
        lbl_t.setStyleSheet("color:white;")
        lbl_t.setAlignment(Qt.AlignmentFlag.AlignCenter)
        lbl_v = QLabel(value)
        lbl_v.setFont(QFont("Arial", 20, QFont.Weight.Bold))
        lbl_v.setStyleSheet("color:white;")
        lbl_v.setAlignment(Qt.AlignmentFlag.AlignCenter)
        lbl_v.setObjectName("val")
        lay.addWidget(lbl_t); lay.addWidget(lbl_v)
        box.setLayout(lay)
        return box

    def _set_metric(self, box, text):
        lbl = box.findChild(QLabel, "val")
        if lbl:
            lbl.setText(text)

    def _find_row_by_tid(self, tid: int) -> int:
        return self.tid_to_row.get(tid, -1)

    # ===== SETUP TAB =====
    def _make_setup_tab(self):
        tab = QWidget()
        lay = QHBoxLayout()

        left     = QWidget()
        left_lay = QVBoxLayout()

        vg  = QGroupBox(" Video")
        vgl = QVBoxLayout()
        self.btn_load    = QPushButton(" Load Video")
        self.btn_load.clicked.connect(self._load_video)
        self.btn_load.setStyleSheet(self._btn_style("#3498db"))
        self.btn_default = QPushButton(" Load Default")
        self.btn_default.clicked.connect(self._load_default)
        self.btn_default.setStyleSheet(self._btn_style("#95a5a6"))
        self.btn_camera = QPushButton(" Connect to Camera!")
        self.btn_camera.setStyleSheet(self._btn_style("#14c741"))
        self.video_info  = QTextEdit()
        self.video_info.setReadOnly(True)
        self.video_info.setMaximumHeight(80)
        vgl.addWidget(self.btn_load); vgl.addWidget(self.btn_default); vgl.addWidget(self.btn_camera) ;vgl.addWidget(self.video_info)
        vg.setLayout(vgl)

        lg  = QGroupBox(" Lane Detection")
        lgl = QVBoxLayout()
        self.btn_detect_lanes = QPushButton(" Detect Lanes")
        self.btn_detect_lanes.clicked.connect(self._detect_lanes)
        self.btn_detect_lanes.setStyleSheet(self._btn_style("#2ecc71", bold=True))
        self.lane_status = QLabel("Not started")
        lgl.addWidget(self.btn_detect_lanes); lgl.addWidget(self.lane_status)
        lg.setLayout(lgl)

        rg  = QGroupBox(" ROI Selection")
        rgl = QVBoxLayout()
        self.roi_status = QLabel(" Not Set")
        self.roi_status.setFont(QFont("Arial", 11, QFont.Weight.Bold))
        self.roi_count  = QLabel("Points: 0/4")
        btn_row = QHBoxLayout()
        btn_clear    = QPushButton("Clear")
        btn_clear.clicked.connect(self._clear_roi)
        btn_clear.setStyleSheet(self._btn_style("#e74c3c", bold=True))
        btn_complete = QPushButton(" Complete")
        btn_complete.clicked.connect(self._complete_roi)
        btn_complete.setStyleSheet(self._btn_style("#2ecc71", bold=True))
        btn_default  = QPushButton(" Default")
        btn_default.clicked.connect(self._load_default_roi)
        btn_default.setStyleSheet(self._btn_style("#3498db", bold=True))
        btn_row.addWidget(btn_clear); btn_row.addWidget(btn_complete); btn_row.addWidget(btn_default)
        self.roi_text = QTextEdit()
        self.roi_text.setReadOnly(True)
        self.roi_text.setMaximumHeight(80)
        rgl.addWidget(self.roi_status); rgl.addWidget(self.roi_count)
        rgl.addLayout(btn_row); rgl.addWidget(self.roi_text)
        rg.setLayout(rgl)

        left_lay.addWidget(vg); left_lay.addWidget(lg); left_lay.addWidget(rg)
        left_lay.addStretch()
        left.setLayout(left_lay)

        right     = QGroupBox(" Frame Preview — Click to add ROI points")
        right_lay = QVBoxLayout()
        self.preview = ClickableLabel()
        self.preview.clicked.connect(self._add_roi_point)
        self.preview.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.preview.setStyleSheet("border:2px solid #3498db; background:#ecf0f1;")
        self.preview.setMinimumSize(900, 550)
        self.preview.setText("No frame loaded")
        right_lay.addWidget(self.preview)
        right.setLayout(right_lay)

        sp = QSplitter(Qt.Orientation.Horizontal)
        sp.addWidget(left); sp.addWidget(right)
        sp.setStretchFactor(0, 1); sp.setStretchFactor(1, 2)
        lay.addWidget(sp)
        tab.setLayout(lay)
        return tab

    # ===== DETECTION TAB =====
    def _make_detection_tab(self):
        tab = QWidget()
        lay = QHBoxLayout()

        left     = QWidget()
        left_lay = QVBoxLayout()
        ctrl = QHBoxLayout()
        self.btn_start = QPushButton(" Start")
        self.btn_start.clicked.connect(self._start)
        self.btn_start.setStyleSheet(self._btn_style("#2ecc71", size=15, bold=True))
        self.btn_stop = QPushButton(" Stop")
        self.btn_stop.clicked.connect(self._stop)
        self.btn_stop.setEnabled(False)
        self.btn_stop.setStyleSheet(self._btn_style("#e74c3c", size=15, bold=True))
        self.sys_status = QLabel(" Ready")

        ctrl.addWidget(self.btn_start)
        ctrl.addWidget(self.btn_stop)
        ctrl.addWidget(self.sys_status)
        ctrl.addSpacing(16)

        # Giới hạn tốc độ theo làn
        lbl_title = QLabel("Speed limit:")
        lbl_title.setStyleSheet("color: #000000; font-size: 12px;")
        ctrl.addWidget(lbl_title)

        for i in range(1, 5):
            limit = LANE_SPEED_LIMITS_UI[i]['max']
            lbl = QLabel(f"Làn {i}: {limit} km/h")
            lbl.setStyleSheet(
                "color: white; font-weight: bold; font-size: 12px;"
                "background: #e74c3c; border-radius: 4px; padding: 3px 10px;"
            )
            ctrl.addWidget(lbl)

        ctrl.addStretch()

        self.video_lbl = QLabel()
        self.video_lbl.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.video_lbl.setStyleSheet("border:3px solid #2ecc71; background:black;")
        self.video_lbl.setMinimumSize(900, 680)
        self.video_lbl.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Expanding)
        self.video_lbl.setText("Press Start")

        left_lay.addLayout(ctrl); left_lay.addWidget(self.video_lbl)
        left.setLayout(left_lay)

        # Right panel — Analytics (khai báo right_lay 1 lần duy nhất)
        right     = QWidget()
        right_lay = QVBoxLayout()

        lbl_analytics = QLabel("<b> Analytics</b>")

        self.chart_speed = MplCanvas(width=5, height=2.8)
        self.chart_lane  = MplCanvas(width=5, height=2.8)
        self._draw_speed_chart([])
        self._draw_time_chart([])
        right_lay.addWidget(self.chart_speed)
        right_lay.addWidget(self.chart_lane)
        right.setLayout(right_lay)

        sp = QSplitter(Qt.Orientation.Horizontal)
        sp.addWidget(left); sp.addWidget(right)
        sp.setStretchFactor(0, 2); sp.setStretchFactor(1, 1)
        lay.addWidget(sp)
        tab.setLayout(lay)
        return tab

    # ===== VIOLATIONS TAB =====
    def _make_violations_tab(self):
        tab = QWidget()
        lay = QHBoxLayout()

        # ---- Left: Export buttons ----
        left     = QWidget()
        left_lay = QVBoxLayout()
        eg  = QGroupBox("📤 Export")
        egl = QVBoxLayout()
        self.btn_export_speed = QPushButton(" Export Speed CSV")
        self.btn_export_speed.clicked.connect(self._export_speed_csv)
        self.btn_export_dist  = QPushButton(" Export Distance CSV")
        self.btn_export_dist.clicked.connect(self._export_dist_csv)
        self.export_status = QLabel("")
        egl.addWidget(self.btn_export_speed)
        egl.addWidget(self.btn_export_dist)
        egl.addWidget(self.export_status)
        eg.setLayout(egl)
        left_lay.addWidget(eg)
        left_lay.addStretch()
        left.setLayout(left_lay)

        # ---- Center: 2 bảng trên/dưới ----
        center     = QWidget()
        center_lay = QVBoxLayout()

        speed_grp = QGroupBox(" Speed Violation Log")
        speed_lay = QVBoxLayout()
        self.vtable = QTableWidget()
        self.vtable.setColumnCount(7)
        self.vtable.setHorizontalHeaderLabels([
            "Time", "Date", "Speed (km/h)", "Over By (km/h)", "Plate", "Color", "Dist Violation"
        ])
        self.vtable.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        self.vtable.cellClicked.connect(self._on_violation_row_clicked)
        speed_lay.addWidget(self.vtable)
        speed_grp.setLayout(speed_lay)

        dist_grp = QGroupBox(" Distance Violation Log")
        dist_lay = QVBoxLayout()
        self.dtable = QTableWidget()
        self.dtable.setColumnCount(6)
        self.dtable.setHorizontalHeaderLabels([
            "Time", "Lane", "ID", "Plate", "Violation Distance (m)", "Safe Dist (m)"
        ])
        self.dtable.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        self.dtable.cellClicked.connect(self._on_dist_violation_row_clicked)
        dist_lay.addWidget(self.dtable)
        dist_grp.setLayout(dist_lay)

        v_splitter = QSplitter(Qt.Orientation.Vertical)
        v_splitter.addWidget(speed_grp)
        v_splitter.addWidget(dist_grp)
        v_splitter.setStretchFactor(0, 1)
        v_splitter.setStretchFactor(1, 1)
        center_lay.addWidget(v_splitter)
        center.setLayout(center_lay)

        sp = QSplitter(Qt.Orientation.Horizontal)
        sp.addWidget(left)
        sp.addWidget(center)
        sp.setStretchFactor(0, 1)
        sp.setStretchFactor(1, 4)
        lay.addWidget(sp)
        tab.setLayout(lay)
        return tab

    # ===== HANDLERS =====
    def _load_video(self):
        path, _ = QFileDialog.getOpenFileName(
            self, "Select Video", "", "Video Files (*.mp4 *.avi *.MOV)"
        )
        if path:
            self._open_video(path)

    def _load_default(self):
        self._open_video(VIDEO_PATH)

    def _open_video(self, path):
        self.video_path = path
        cap = cv2.VideoCapture(path)
        if not cap.isOpened():
            self.video_info.setText(" Cannot open video")
            return
        self.fps = cap.get(cv2.CAP_PROP_FPS)
        ret, frame = cap.read()
        cap.release()
        if ret:
            self.first_frame = frame.copy()
            self._show_frame(self.preview, frame)
        self.video_info.setText(f"Path: {Path(path).name}\nFPS: {self.fps:.2f}")

    def _load_models(self):
        if not self.lane_model:
            self.lane_model = YOLO(LANE_MODEL_PATH)
        if not self.car_model:
            self.car_model = YOLO(CAR_MODEL_PATH)

    def _detect_lanes(self):
        if self.first_frame is None:
            self.lane_status.setText(" Load video first!")
            return
        self._load_models()
        self.lane_status.setText(" Detecting...")
        QApplication.processEvents()
        frame_rgb = cv2.cvtColor(self.first_frame, cv2.COLOR_BGR2RGB)
        self.solid_coeffs, self.dotted_coeffs, self.all_lines, self.vp = \
            detect_lanes(self.lane_model, frame_rgb)
        preview = draw_lanes_on_frame(
            self.first_frame, self.solid_coeffs, self.dotted_coeffs,
            self.all_lines, self.vp
        )
        self._show_frame(self.preview, preview)
        n_lanes = max(0, len(self.all_lines) - 1)
        self.lane_status.setText(f" {n_lanes} lanes detected")

    def _add_roi_point(self, x, y):
        if self.first_frame is None or len(self.roi_pts) >= 4:
            return
        px = self.preview.pixmap()
        if not px:
            return
        sx = self.first_frame.shape[1] / px.width()
        sy = self.first_frame.shape[0] / px.height()
        ox = (self.preview.width()  - px.width())  / 2
        oy = (self.preview.height() - px.height()) / 2
        rx, ry = int((x - ox) * sx), int((y - oy) * sy)
        self.roi_pts.append([rx, ry])
        self.roi_count.setText(f"Points: {len(self.roi_pts)}/4")
        if len(self.roi_pts) == 4:
            self.roi_status.setText(" Complete (4/4)")
        self._redraw_roi()

    def _redraw_roi(self):
        if self.first_frame is None:
            return
        frame = draw_lanes_on_frame(
            self.first_frame, self.solid_coeffs, self.dotted_coeffs,
            self.all_lines, self.vp
        ) if self.solid_coeffs else self.first_frame.copy()

        for i, pt in enumerate(self.roi_pts):
            cv2.circle(frame, tuple(pt), 8, (0, 255, 0), -1)
            cv2.putText(frame, str(i + 1), (pt[0] + 10, pt[1] - 10),
                        cv2.FONT_HERSHEY_SIMPLEX, 0.8, (0, 255, 0), 2)
        if len(self.roi_pts) > 1:
            pts = np.array(self.roi_pts, np.int32)
            cv2.polylines(frame, [pts], len(self.roi_pts) >= 4, (0, 255, 0), 3)
        if len(self.roi_pts) >= 4:
            overlay = frame.copy()
            cv2.fillPoly(overlay, [np.array(self.roi_pts, np.int32)], (0, 255, 0))
            frame = cv2.addWeighted(overlay, 0.3, frame, 0.7, 0)
        self._show_frame(self.preview, frame)

    def _clear_roi(self):
        self.roi_pts = []
        self.roi_status.setText(" Not Set")
        self.roi_count.setText("Points: 0/4")
        self.roi_text.clear()
        self._redraw_roi()

    def _complete_roi(self):
        if len(self.roi_pts) != 4:
            self.roi_status.setText(" Need 4 points")
            return
        self.roi_status.setText(" Ready")
        self.roi_text.setText("\n".join(f"P{i+1}: ({p[0]}, {p[1]})" for i, p in enumerate(self.roi_pts)))

    def _load_default_roi(self):
        self.roi_pts = [list(p) for p in DEFAULT_ROI]
        self.roi_status.setText(" Default ROI loaded")
        self.roi_count.setText("Points: 4/4")
        self.roi_text.setText("\n".join(f"P{i+1}: ({p[0]}, {p[1]})" for i, p in enumerate(self.roi_pts)))
        self._redraw_roi()

    def _start(self):
        if self.first_frame is None or len(self.roi_pts) != 4:
            self.sys_status.setText(" Complete setup first!")
            return
        self._load_models()

        for q in [self.plate_queue, self.color_queue]:
            while not q.empty():
                q.get_nowait()
        self.tid_to_row.clear()
        self.tid_to_frame_path.clear()
        self.tid_to_dist_row.clear()
        self.tid_to_dist_frame_path.clear()  

        self.plate_thread = PlateOCRThread(PLATE_MODEL_PATH, self.plate_queue)
        self.plate_thread.plate_detected.connect(self._on_plate_detected)
        self.plate_thread.start()

        self.color_thread = ColorDetectThread(self.color_queue)
        self.color_thread.color_detected.connect(self._on_color_detected)
        self.color_thread.start()

        all_lines_roi = filter_coeffs_in_roi(
            self.all_lines, self.roi_pts, self.first_frame.shape[0]
        )
        self.estimator = SpeedEstimator(
            self.roi_pts, self.fps, LANE_SPEED_LIMITS,
            crops_output_dir=CROPS_OUTPUT_DIR,
            plate_queue=self.plate_queue,
            color_queue=self.color_queue,
        )
        self.thread.setup(
            self.video_path, self.car_model, self.estimator,
            self.solid_coeffs, self.dotted_coeffs, all_lines_roi,
            self.vp, self.roi_pts
        )
        self.thread.start()
        self.btn_start.setEnabled(False)
        self.btn_stop.setEnabled(True)
        self.sys_status.setText(" Processing...")

    def _stop(self):
        for t in [self.thread, self.plate_thread, self.color_thread]:
            if t:
                t.stop()
                t.wait(3000)
        self.btn_start.setEnabled(True)
        self.btn_stop.setEnabled(False)

    # ===== THREAD CALLBACKS =====
    def _on_frame(self, frame):
        self._show_frame(self.video_lbl, frame)

    def _on_stats(self, stats):
        self._set_metric(self.m_vehicles,   str(stats['vehicle_count']))
        self._set_metric(self.m_violations, str(stats['violations_count']))
        rate = stats['violations_count'] / max(stats['vehicle_count'], 1) * 100
        self._set_metric(self.m_rate,      f"{rate:.1f}%")
        self._set_metric(self.m_tracking,   str(stats['current_tracking']))
        self.sys_status.setText(" Processing...")
        if self.estimator and stats.get('frame_count', 0) % 5 == 0 and self.estimator.all_speeds:
            self._draw_speed_chart(self.estimator.all_speeds)
            self._draw_time_chart(self.estimator.violations)

    def _on_violation(self, v):
        """Thêm row mới vào bảng Speed Violations, lưu mapping tid → row và frame path."""
        row = self.vtable.rowCount()
        self.vtable.insertRow(row)
        self.tid_to_row[v['ID']] = row

        # Tạo path frame khớp đúng convention trong speed_estimation.py
        self.tid_to_frame_path[v['ID']] = v.get('Frame Path', '')
        
        insert_speed_violation(v)

        dist_val = "True" if v.get('Dist Violation', False) else "False"
        for col, val in enumerate([
            v['Time'], v['Date'],
            str(v['Speed (km/h)']),
            str(v.get('Over By (km/h)', '')),
            v.get('Plate', 'Detecting...'),
            v.get('Color', 'Detecting...'),
            dist_val,
        ]):
            item = QTableWidgetItem(val)
            if col == self.COL_OVER:
                try:
                    if int(val) <= 5:
                        item.setForeground(Qt.GlobalColor.black)
                    elif int(val)<=10:
                        item.setForeground(Qt.GlobalColor.yellow)
                    else:
                        item.setForeground(Qt.GlobalColor.red)
                except:
                    pass
            if col == self.COL_DIST and val == "True":
                item.setForeground(Qt.GlobalColor.red)
            self.vtable.setItem(row, col, item)

    def _on_dist_violation(self, dv):
        row = self.dtable.rowCount()
        self.dtable.insertRow(row)

        self.tid_to_dist_row[dv['ID_Behind']] = row
        self.tid_to_dist_frame_path[dv['ID_Behind']] = dv.get('Frame Path', '') 

        insert_distance_violation(dv)
        update_dist_violation(dv['ID_Behind'])

        for col, val in enumerate([
            dv['Time'],
            dv['Lane'],
            str(dv['ID_Behind']),
            dv.get('Plate', 'Detecting...'),
            str(dv['Distance (m)']),
            str(dv['Safe_Dist (m)']),
        ]):
            self.dtable.setItem(row, col, QTableWidgetItem(val))

        row_v = self._find_row_by_tid(dv['ID_Behind'])
        if row_v != -1:
            item = QTableWidgetItem("True")
            item.setForeground(Qt.GlobalColor.red)
            self.vtable.setItem(row_v, self.COL_DIST, item)

        if self.estimator:
            self._draw_time_chart(self.estimator.violations)

    def _on_plate_detected(self, tid: int, plate_text: str):
        if self.estimator:
            for v in self.estimator.violations:
                if v['ID'] == tid:
                    v['Plate'] = plate_text
                    break

        update_plate(tid, plate_text)
        update_dist_plate(tid, plate_text)

        # Cập nhật vtable
        row = self._find_row_by_tid(tid)
        if row != -1:
            self.vtable.setItem(row, self.COL_PLATE, QTableWidgetItem(plate_text))

        # Cập nhật dtable — tìm row có cột ID (col 2) == tid
        for row in range(self.dtable.rowCount()):
            item = self.dtable.item(row, 2)
            if item and item.text() == str(tid):
                self.dtable.setItem(row, 3, QTableWidgetItem(plate_text))
                break
            

    def _on_color_detected(self, tid: int, color_name: str):
        if self.estimator:
            for v in self.estimator.violations:
                if v['ID'] == tid:
                    v['Color'] = color_name
                    break
                
        update_color(tid, color_name)
        
        row = self._find_row_by_tid(tid)
        if row != -1:
            self.vtable.setItem(row, self.COL_COLOR, QTableWidgetItem(color_name))

    def _on_violation_row_clicked(self, row: int, col: int):
        """Click row → chờ file sẵn rồi mở popup, không blocking UI."""
        tid = next((t for t, r in self.tid_to_row.items() if r == row), None)
        if tid is None:
            return
        path = self.tid_to_frame_path.get(tid)
        if not path:
            return
        self._open_popup_when_ready(tid, path, retries=20)

    def _open_popup_when_ready(self, tid: int, path: str, retries: int):
        """Retry tối đa 20 lần × 150ms = 3 giây chờ file được write xong."""
        from PyQt6.QtCore import QTimer
        if Path(path).exists():
            self._open_popup(tid, path)
        elif retries > 0:
            QTimer.singleShot(150, lambda: self._open_popup_when_ready(tid, path, retries - 1))
        else:
            self._show_popup_msg(f"Image not found:\n{path}")

    def _open_popup(self, tid: int, path: str):
        """Mở cửa sổ popup hiển thị full frame có bbox."""
        img = cv2.imread(path)
        if img is None:
            self._show_popup_msg("Cannot load image")
            return

        rgb       = cv2.cvtColor(img, cv2.COLOR_BGR2RGB)
        h, w, ch  = rgb.shape
        qt_img    = QImage(rgb.data, w, h, ch * w, QImage.Format.Format_RGB888)
        px        = QPixmap.fromImage(qt_img)

        popup = QWidget()
        popup.setWindowTitle(f"Vehicle #{tid} — {Path(path).name}")
        popup.setGeometry(300, 150, 900, 600)
        lay = QVBoxLayout()
        lbl = QLabel()
        lbl.setAlignment(Qt.AlignmentFlag.AlignCenter)
        lbl.setPixmap(px.scaled(
            880, 560,
            Qt.AspectRatioMode.KeepAspectRatio,
            Qt.TransformationMode.SmoothTransformation
        ))
        lay.addWidget(lbl)
        popup.setLayout(lay)
        popup.show()
        self._popup = popup  # giữ reference tránh bị garbage collected
        
    def _on_dist_violation_row_clicked(self, row: int, col: int):
        # Lấy tid từ cột ID (col 2)
        item = self.dtable.item(row, 2)
        if item is None:
            return
        try:
            tid = int(item.text())
        except ValueError:
            return

        # Lấy thêm thông tin để làm title popup
        time_item = self.dtable.item(row, 0)
        lane_item = self.dtable.item(row, 1)
        dist_item = self.dtable.item(row, 4)
        time_str  = time_item.text() if time_item else ''
        lane_str  = lane_item.text() if lane_item else ''
        dist_str  = dist_item.text() if dist_item else ''

        title = f"Khoảng cách — Xe #{tid} | {lane_str} | {dist_str}m | {time_str}"

        path = self.tid_to_dist_frame_path.get(tid)
        if not path:
            self._show_popup_msg("Không có ảnh cho vi phạm này")
            return

        self._open_popup_when_ready(tid, path, retries=20, title=title)


    def _open_popup_when_ready(self, tid: int, path: str, retries: int, title: str = None):
        from PyQt6.QtCore import QTimer
        if Path(path).exists():
            self._open_popup(tid, path, title=title)
        elif retries > 0:
            QTimer.singleShot(150, lambda: self._open_popup_when_ready(tid, path, retries - 1, title=title))
        else:
            self._show_popup_msg(f"Image not found:\n{path}")


    def _open_popup(self, tid: int, path: str, title: str = None):
        img = cv2.imread(path)
        if img is None:
            self._show_popup_msg("Cannot load image")
            return

        rgb      = cv2.cvtColor(img, cv2.COLOR_BGR2RGB)
        h, w, ch = rgb.shape
        qt_img   = QImage(rgb.data, w, h, ch * w, QImage.Format.Format_RGB888)
        px       = QPixmap.fromImage(qt_img)

        window_title = title if title else f"Vehicle #{tid} — {Path(path).name}"

        popup = QWidget()
        popup.setWindowTitle(window_title)
        popup.setGeometry(300, 150, 900, 600)
        lay = QVBoxLayout()
        lbl = QLabel()
        lbl.setAlignment(Qt.AlignmentFlag.AlignCenter)
        lbl.setPixmap(px.scaled(
            880, 560,
            Qt.AspectRatioMode.KeepAspectRatio,
            Qt.TransformationMode.SmoothTransformation
        ))
        lay.addWidget(lbl)
        popup.setLayout(lay)
        popup.show()
        self._popup = popup

    def _show_popup_msg(self, msg: str):
        """Popup thông báo lỗi nhỏ."""
        popup = QWidget()
        popup.setWindowTitle("Notice")
        popup.setGeometry(400, 300, 350, 80)
        lay = QVBoxLayout()
        lbl = QLabel(msg)
        lbl.setAlignment(Qt.AlignmentFlag.AlignCenter)
        lay.addWidget(lbl)
        popup.setLayout(lay)
        popup.show()
        self._popup = popup

    def _on_finished(self):
        self.btn_start.setEnabled(True)
        self.btn_stop.setEnabled(False)
        self._export_excel_auto()

    # ===== CHARTS =====
    def _draw_speed_chart(self, all_speeds):
        ax = self.chart_speed.axes
        ax.clear()
        if not all_speeds:
            ax.text(0.5, 0.5, 'No data', ha='center', va='center', transform=ax.transAxes)
        else:
            total    = len(all_speeds)
            n_vio    = len(self.estimator.violations) if self.estimator else 0
            n_normal = total - n_vio
            vals = [v for v in [n_normal, n_vio] if v > 0]
            keys = [k for k, v in [("", n_normal), ("Vi phạm \ntốc độ", n_vio)] if v > 0]
            ax.pie(vals, labels=keys, autopct='%1.1f%%',
                   colors=['#2ecc71', '#e74c3c'])
            ax.text(0.5, 1.02, f"Tỉ lệ vi phạm",
                    transform=ax.transAxes,
                    ha='center', va='bottom', fontsize=10, fontweight='bold')
        self.chart_speed.draw()

    def _draw_time_chart(self, violations):
        ax = self.chart_lane.axes
        ax.clear()

        if not self.estimator:
            ax.text(0.5, 0.5, 'No data', ha='center', va='center', transform=ax.transAxes)
            self.chart_lane.draw()
            return

        # Tổng xe đã đếm
        total = len(self.estimator.saved_ids)
        # Số xe vi phạm khoảng cách (unique ID_Behind)
        n_dist = len({dv['ID_Behind'] for dv in self.estimator.distance_violations})
        n_normal = total - n_dist

        if total == 0:
            ax.text(0.5, 0.5, 'No data', ha='center', va='center', transform=ax.transAxes)
            self.chart_lane.draw()
            return

        vals = [v for v in [n_normal, n_dist] if v > 0]
        keys = [k for k, v in [("", n_normal), ("Vi phạm \nkhoảng cách", n_dist)] if v > 0]

        ax.pie(vals, labels=keys, autopct='%1.1f%%',
            colors=['#2ecc71', '#f39c12'])
        ax.text(0.5, 1.02, "Tỉ lệ vi phạm khoảng cách",
                transform=ax.transAxes,
                ha='center', va='bottom', fontsize=10, fontweight='bold')
        self.chart_lane.draw()

    # ===== EXPORT =====
    def _export_speed_csv(self):
        if not self.estimator or not self.estimator.violations:
            self.export_status.setText(" No speed data to export")
            return
        path, _ = QFileDialog.getSaveFileName(
            self, "Save Speed CSV",
            f"violations_speed_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv",
            "CSV Files (*.csv)"
        )
        if path:
            self._build_speed_df().to_csv(path, index=False)
            self.export_status.setText(f" Exported {len(self.estimator.violations)} rows")

    def _export_dist_csv(self):
        if not self.estimator or not self.estimator.distance_violations:
            self.export_status.setText(" No distance data to export")
            return
        path, _ = QFileDialog.getSaveFileName(
            self, "Save Distance CSV",
            f"violations_distance_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv",
            "CSV Files (*.csv)"
        )
        if path:
            self._build_dist_df().to_csv(path, index=False)
            self.export_status.setText(f" Exported {len(self.estimator.distance_violations)} rows")

    def _export_excel_auto(self):
        if not self.estimator:
            return
        Path(EXCEL_OUTPUT_DIR).mkdir(parents=True, exist_ok=True)
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")

        from openpyxl.styles import Font, PatternFill, Alignment

        def _write_sheet(df, path, sheet_name):
            if df.empty:
                return
            with pd.ExcelWriter(path, engine='openpyxl') as writer:
                df.to_excel(writer, index=False, sheet_name=sheet_name)
                ws = writer.sheets[sheet_name]
                header_fill = PatternFill("solid", fgColor="2C3E50")
                header_font = Font(bold=True, color="FFFFFF", name="Arial", size=13)
                for cell in ws[1]:
                    cell.fill      = header_fill
                    cell.font      = header_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                for col in ws.columns:
                    max_len = max(len(str(cell.value or "")) for cell in col)
                    ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)
                alt_fill = PatternFill("solid", fgColor="EBF5FB")
                for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
                    if row_idx % 2 == 0:
                        for cell in row:
                            cell.fill = alt_fill

        if self.estimator.violations:
            try:
                _write_sheet(
                    self._build_speed_df(),
                    f"{EXCEL_OUTPUT_DIR}/violations_speed_{ts}.xlsx",
                    "Speed Violations"
                )
            except Exception as e:
                print(f" Speed Excel error: {e}")

        if self.estimator.distance_violations:
            try:
                _write_sheet(
                    self._build_dist_df(),
                    f"{EXCEL_OUTPUT_DIR}/violations_distance_{ts}.xlsx",
                    "Distance Violations"
                )
            except Exception as e:
                print(f" Distance Excel error: {e}")

        self.export_status.setText(f" Auto-saved to {EXCEL_OUTPUT_DIR}/")

    def _build_speed_df(self):
        if not self.estimator or not self.estimator.violations:
            return pd.DataFrame()
        cols = ['Time', 'Date', 'Speed (km/h)', 'Plate', 'Color', 'Dist Violation']
        return pd.DataFrame(self.estimator.violations)[cols]

    def _build_dist_df(self):
        if not self.estimator or not self.estimator.distance_violations:
            return pd.DataFrame()
        cols = ['Time', 'Date', 'ID_Behind',
                'Distance (m)', 'Safe_Dist (m)', 'Speed (km/h)', 'Lane', 'Plate']
        return pd.DataFrame(self.estimator.distance_violations)[cols]

    # ===== HELPERS =====
    def _show_frame(self, label, frame):
        rgb  = cv2.cvtColor(frame, cv2.COLOR_BGR2RGB)
        h, w, ch = rgb.shape
        img  = QImage(rgb.data, w, h, ch * w, QImage.Format.Format_RGB888)
        px   = QPixmap.fromImage(img)
        label.setPixmap(px.scaled(
            label.size(),
            Qt.AspectRatioMode.KeepAspectRatio,
            Qt.TransformationMode.SmoothTransformation
        ))

    def _btn_style(self, color, size=10, bold=False):
        w = "bold" if bold else "normal"
        return (
            f"QPushButton {{ background-color:{color}; color:white; padding:10px;"
            f" border-radius:5px; font-size:{size}px; font-weight:{w}; }}"
            f"QPushButton:hover {{ background-color:{color}cc; }}"
        )


# ===== ENTRY POINT =====
def main():
    app = QApplication(sys.argv)
    app.setStyle('Fusion')
    win = TrafficApp()
    win.show()
    sys.exit(app.exec())


if __name__ == "__main__":
    main()